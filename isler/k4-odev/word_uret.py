#!/usr/bin/env python3
"""odev.md dosyasından odev.docx üretir (python-docx).
Times New Roman 12, 1,5 satır aralığı, iki yana yaslı, başlık sayfası, içindekiler,
sayfa numarası, biçimli tablolar, kaynakçada asılı girinti, tıklanabilir bağlantılar."""
import re, pathlib
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import docx.opc.constants as C

ROOT = pathlib.Path(__file__).parent
md = (ROOT / "odev.md").read_text(encoding="utf-8").splitlines()
FONT = "Times New Roman"

doc = Document()
sec = doc.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, side, Cm(2.5))

def set_font(style, size, bold=False):
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts"); rpr.append(rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), FONT)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rfonts.get(qn(a)) is not None:
            del rfonts.attrib[qn(a)]

normal = doc.styles["Normal"]
set_font(normal, 12)
normal.paragraph_format.line_spacing = 1.5
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
for name, size in (("Heading 1", 14), ("Heading 2", 13), ("Heading 3", 12)):
    st = doc.styles[name]
    set_font(st, size, bold=True)
    st.paragraph_format.space_before = Pt(12)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.5
    st.paragraph_format.keep_with_next = True
    st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

def add_hyperlink(par, url, text):
    part = par.part
    r_id = part.relate_to(url, C.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), r_id)
    r = OxmlElement("w:r"); rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        fonts.set(qn(a), FONT)
    rpr.append(fonts)
    col = OxmlElement("w:color"); col.set(qn("w:val"), "0563C1"); rpr.append(col)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rpr.append(u)
    r.append(rpr)
    t = OxmlElement("w:t"); t.text = text; t.set(qn("xml:space"), "preserve"); r.append(t)
    h.append(r); par._p.append(h)

TOKEN = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|https?://\S+)")
def add_inline(par, text, size=None):
    text = text.replace("\\*", "\u0001")
    for part in TOKEN.split(text):
        part = part.replace("\u0001", "*") if not (part.startswith("*") and part.endswith("*")) else part
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = par.add_run(part[2:-2]); run.bold = True
        elif part.startswith("http"):
            url = part.rstrip(".,;)")
            tail = part[len(url):]
            add_hyperlink(par, url, url)
            if tail:
                run = par.add_run(tail)
            continue
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = par.add_run(part[1:-1].replace("\u0001", "*")); run.italic = True
        else:
            run = par.add_run(part)
        if size:
            run.font.size = Pt(size)

def add_field(par, instr):
    r = par.add_run()
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin")
    i = OxmlElement("w:instrText"); i.set(qn("xml:space"), "preserve"); i.text = instr
    s = OxmlElement("w:fldChar"); s.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t"); t.text = "İçindekiler tablosunu görmek için sağ tıklayıp 'Alanı Güncelleştir' seçin."
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
    r._r.append(b); r._r.append(i); r._r.append(s); r._r.append(t); r._r.append(e)

def shade(cell, hexcolor):
    tcpr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hexcolor)
    tcpr.append(sh)

def add_table(rows):
    header = [c.strip() for c in rows[0].strip("|").split("|")]
    body = [[c.strip() for c in r.strip("|").split("|")] for r in rows[2:]]
    t = doc.add_table(rows=1 + len(body), cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate([header] + body):
        for j, txt in enumerate(row):
            cell = t.cell(i, j)
            p = cell.paragraphs[0]
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, txt, size=10)
            if i == 0:
                for r in p.runs:
                    r.bold = True
                shade(cell, "D9D9D9")
    # başlık satırı her sayfada tekrar etsin
    trpr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trpr.append(th)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

# ---------- Başlık sayfası ----------
title = md[0].lstrip("# ").strip()
for _ in range(5):
    doc.add_paragraph()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run(title.upper()); r.bold = True; r.font.size = Pt(16)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Yüksek Lisans Dönem Ödevi").italic = True
for _ in range(6):
    doc.add_paragraph()
for label in ("Öğrenci Adı Soyadı", "Öğrenci Numarası", "Ders", "Dersin Sorumlusu", "Teslim Tarihi"):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.left_indent = Cm(3)
    r = p.add_run(f"{label}: "); r.bold = True
    p.add_run("_" * 40)
doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

# ---------- İçindekiler ----------
p = doc.add_paragraph(); r = p.add_run("İçindekiler"); r.bold = True; r.font.size = Pt(14)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_field(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u')
# ana metin yeni bölümde başlasın, sayfa numarası 1'den
new_sec = doc.add_section(WD_SECTION.NEW_PAGE)
pg = new_sec._sectPr
pn = OxmlElement("w:pgNumType"); pn.set(qn("w:start"), "1"); pg.append(pn)
new_sec.footer.is_linked_to_previous = False
fp = new_sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run()
for typ, txt in (("begin", None), ("instr", "PAGE"), ("separate", None), ("text", "1"), ("end", None)):
    if typ == "instr":
        el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = txt
    elif typ == "text":
        el = OxmlElement("w:t"); el.text = txt
    else:
        el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), typ)
    r._r.append(el)
r.font.name = FONT; r.font.size = Pt(11)

# ---------- Gövde ----------
mode = "body"
i = 1
while i < len(md):
    line = md[i]
    s = line.strip()
    if not s:
        i += 1; continue
    if s.startswith("|"):
        rows = []
        while i < len(md) and md[i].strip().startswith("|"):
            rows.append(md[i].strip()); i += 1
        add_table(rows); continue
    if s.startswith("## "):
        head = s[3:]
        if head.startswith("Kaynakça") or head.startswith("Ek"):
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        mode = "refs" if head.startswith("Kaynakça") else ("ek" if head.startswith("Ek") else "body")
        doc.add_heading(head, level=1)
        i += 1; continue
    if s.startswith("### "):
        doc.add_heading(s[4:], level=2); i += 1; continue
    if s.startswith("- "):
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.line_spacing = 1.5
        add_inline(p, s[2:]); i += 1; continue
    m = re.match(r"^(\d+)\. (.*)", s)
    if m:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.9)
        p.paragraph_format.first_line_indent = Cm(-0.6)
        p.add_run(m.group(1) + ". ")
        add_inline(p, m.group(2)); i += 1; continue
    p = doc.add_paragraph()
    if mode == "refs":
        p.paragraph_format.left_indent = Cm(1.27)
        p.paragraph_format.first_line_indent = Cm(-1.27)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    elif s.startswith("**Tablo") or s.startswith("Kaynak:"):
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.keep_with_next = s.startswith("**Tablo")
        if s.startswith("Kaynak:"):
            p.paragraph_format.space_after = Pt(12)
    add_inline(p, s, size=10 if s.startswith("Kaynak:") else None)
    i += 1

# Word açılışta alanları güncellesin (içindekiler)
settings = doc.settings.element
upd = OxmlElement("w:updateFields"); upd.set(qn("w:val"), "true"); settings.append(upd)

out = ROOT / "odev.docx"
doc.save(out)
print("yazıldı:", out)
